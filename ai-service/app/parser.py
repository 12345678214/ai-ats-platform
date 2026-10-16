"""Resume parsing: text extraction (PDF/DOCX) + structured extraction (skills, education, experience)."""
import io
import re
from datetime import date

from . import taxonomy

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
MONTH_RE = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
DATE_TOKEN = rf"(?:{MONTH_RE}\.?\s+)?(?:19|20)\d{{2}}"
RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:-|–|—|to)\s*(?P<end>{DATE_TOKEN}|present|current|now|ongoing|till date)",
    re.IGNORECASE,
)
SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "objective", "about me", "about"],
    "experience": ["experience", "work experience", "professional experience", "employment", "employment history", "work history", "internships", "internship"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "education": ["education", "academic background", "academics", "qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tech stack", "key skills"],
    "achievements": ["achievements", "awards", "certifications", "certificates", "accomplishments", "honors"],
}
HEADING_LOOKUP = {a: k for k, v in SECTION_ALIASES.items() for a in v}

DEGREE_LEVELS = [
    (4, r"ph\.?\s?d|doctorate|doctor of philosophy"),
    (3, r"master(?:'s|s)?(?: of)?|m\.?\s?tech|m\.?\s?sc|m\.?\s?s\b|m\.?\s?e\b|mba|mca|m\.?\s?eng"),
    (2, r"bachelor(?:'s|s)?(?: of)?|b\.?\s?tech|b\.?\s?e\b|b\.?\s?sc|b\.?\s?s\b|bca|b\.?\s?eng|b\.?\s?a\b|undergraduate degree"),
    (1, r"diploma|associate(?:'s)? degree|polytechnic"),
]
LEVEL_NAMES = {4: "PhD", 3: "Master's", 2: "Bachelor's", 1: "Diploma", 0: "Unknown"}
INSTITUTION_RE = re.compile(r"\b(university|college|institute|school|academy|polytechnic)\b", re.IGNORECASE)


# ───────────────────────── text extraction ─────────────────────────
def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return _pdf_text(data)
    if name.endswith(".docx") or data[:2] == b"PK":
        return _docx_text(data)
    raise ValueError("Unsupported file type. Upload a PDF or DOCX.")


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    pages = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > 15:
                raise ValueError("Resume has too many pages (limit 15)")
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except ValueError:
        raise
    except Exception:
        pages = []
    text = "\n".join(pages).strip()
    if not text:  # fall back to pypdf for odd encodings
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages[:15]).strip()
    if not text:
        raise ValueError("No readable text found. Scanned/image-only PDFs are not supported yet.")
    return _clean(text)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    text = "\n".join(lines).strip()
    if not text:
        raise ValueError("The document contains no readable text")
    return _clean(text)


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\u2022", "•").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ───────────────────────── sections ─────────────────────────
def split_sections(text: str) -> dict:
    sections, current = {"header": []}, "header"
    for line in text.splitlines():
        key = HEADING_LOOKUP.get(re.sub(r"[^a-z& ]", "", line.strip().lower()).strip())
        if key and len(line.strip()) < 40:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


# ───────────────────────── experience ─────────────────────────
def _to_month_index(token: str, is_end: bool):
    token = token.strip().lower()
    if token in {"present", "current", "now", "ongoing", "till date"}:
        today = date.today()
        return today.year * 12 + today.month
    year = int(re.search(r"(?:19|20)\d{2}", token).group())
    month_match = re.match(MONTH_RE, token)
    month = MONTHS[month_match.group()[:3]] if month_match else (12 if is_end else 1)
    return year * 12 + month


def date_ranges(text: str):
    ranges = []
    for m in RANGE_RE.finditer(text):
        start, end = _to_month_index(m.group("start"), False), _to_month_index(m.group("end"), True)
        if end >= start and end - start <= 12 * 40:
            ranges.append((start, end, m.group(0), m.start()))
    return ranges


def years_from_ranges(ranges) -> float:
    """Union of month ranges so overlapping jobs are not double counted."""
    spans = sorted((s, e + 1) for s, e, *_ in ranges)  # inclusive months
    total, cur_s, cur_e = 0, None, None
    for s, e in spans:
        if cur_e is None or s > cur_e:
            if cur_e is not None:
                total += cur_e - cur_s
            cur_s, cur_e = s, e
        else:
            cur_e = max(cur_e, e)
    if cur_e is not None:
        total += cur_e - cur_s
    return round(total / 12, 1)


def explicit_years(text: str) -> float:
    vals = [float(x) for x in re.findall(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs)\b(?:\s+of)?(?:\s+\w+){0,3}?\s+experience", text, re.IGNORECASE)]
    return max(vals) if vals else 0.0


def parse_experience(sections: dict, text: str):
    source = sections.get("experience") or ""
    entries = []
    lines = source.splitlines()
    for i, line in enumerate(lines):
        for _, _, raw, pos in date_ranges(line):
            head = (line[:pos] + line[pos + len(raw):]).strip(" |-–—,•")
            nxt = next((l.strip(" •-") for l in lines[i + 1:i + 3] if l.strip() and not l.strip().startswith(("•", "-"))), "")
            prev = lines[i - 1].strip() if i > 0 and lines[i - 1].strip() and not lines[i - 1].strip().startswith("•") else ""
            title, company = (nxt, head) if head and nxt else (head or prev, "")
            if title or company:
                entries.append({"title": title[:120], "company": company[:120], "period": raw.strip(), "description": ""})
    return entries[:10]


def compute_years(sections: dict, text: str) -> float:
    explicit = explicit_years(text)
    exp_text = sections.get("experience", "")
    ranges = date_ranges(exp_text)
    computed = years_from_ranges(ranges) if ranges else 0.0
    return round(max(explicit, computed), 1)


# ───────────────────────── education ─────────────────────────
def degree_level(text: str) -> int:
    best = 0
    for level, pattern in DEGREE_LEVELS:
        if re.search(rf"(?<![A-Za-z]){pattern}(?![A-Za-z])", text, re.IGNORECASE):
            best = max(best, level)
    return best


def parse_education(sections: dict):
    source = sections.get("education", "")
    lines = [l.strip() for l in source.splitlines() if l.strip()]
    entries = []
    for i, line in enumerate(lines):
        window = lines[max(0, i - 1): i + 2]
        institution = next((l for l in [line] + window if INSTITUTION_RE.search(l)), "")
        has_degree = degree_level(line) > 0 or re.search(r"\b(bachelor|master|diploma|b\.?tech|m\.?tech)\b", line, re.IGNORECASE)
        if not has_degree:
            continue
        years = re.findall(r"(?:19|20)\d{2}", " ".join(window))
        field = ""
        m = re.search(r"\b(?:in|of)\s+([A-Z][A-Za-z &/]+?)(?:\s*[,|(\d]|$)", line)
        if m:
            field = m.group(1).strip()
        degree = re.sub(r"\s*(?:\(?(?:19|20)\d{2}.*)$", "", line).strip(" ,|-")
        inst_clean = re.sub(r"\s*(?:\(?[A-Za-z]{3,9}\.?\s*)?(?:19|20)\d{2}.*$", "", institution).strip(" ,|-") if institution != line else ""
        entries.append({
            "degree": degree[:120], "field": field[:120], "institution": inst_clean[:160],
            "year": years[-1] if years else "",
        })
    # an institution line like "Teerthanker Mahaveer University, Moradabad Aug 2022 – Jun 2026"
    # followed by "Bachelor of ..." is the typical layout; merge institution into the entry.
    for e in entries:
        if not e["institution"]:
            for l in lines:
                if INSTITUTION_RE.search(l):
                    e["institution"] = re.sub(r"\s*(?:[A-Za-z]{3,9}\.?\s*)?(?:19|20)\d{2}.*$", "", l).strip(" ,|-")[:160]
                    break
    return entries[:5]


# ───────────────────────── skills / contact ─────────────────────────
def extract_skills(text: str):
    found = taxonomy.find_skills(text)
    found.sort(key=lambda x: x[1])
    return [{"name": s.name, "category": s.category} for s, _ in found]


def extract_contact(text: str):
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)\d{3,5}[\s-]?\d{3,5}", text[:600])
    links = re.findall(r"https?://[^\s)|,]+", text)
    return {
        "email": email.group().strip(".") if email else None,
        "phone": phone.group().strip() if phone else None,
        "links": links[:6],
    }


def parse_resume(data: bytes, filename: str) -> dict:
    text = extract_text(data, filename)
    sections = split_sections(text)
    education = parse_education(sections)
    level = max([degree_level(e["degree"]) for e in education] + [degree_level(sections.get("education", ""))])
    return {
        "text": text,
        "skills": extract_skills(text),
        "education": education,
        "experience": parse_experience(sections, text),
        "yearsExperience": compute_years(sections, text),
        "highestDegree": LEVEL_NAMES[level] if level else None,
        "contact": extract_contact(text),
        "sections": [k for k, v in sections.items() if v and k != "header"],
        "wordCount": len(text.split()),
    }
